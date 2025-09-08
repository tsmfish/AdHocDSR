"""
references:
    - https://datatracker.ietf.org/doc/html/rfc4728#:~:text=All%20aspects%20of%20the%20protocol,the%20routes%20currently%20in%20use.
    reference for constants see section: Protocol Constants and Configuration Variables

    - https://www.ietf.org/rfc/rfc4728.txt Raw text RFC
    - https://www.researchgate.net/figure/Overall-Basic-Operation-of-the-DSR-Protocol_fig3_267375905 - simulation sample with anisotropy network
"""

import adhoc_nodes
import random
import cv2
import numpy
import openpyxl


from adhoc_pack import RReqPack, DataPack


# ---------------------------------------------------------------------------------------------------------------
#  Class Ad hoc net
# --------------------------------------------------------------------------------------------------------------
class AdHocNet:
    def __init__(self):
        # base parameters   --------------------------------------------------
        self._plase_x_size = 1700
        self._plase_y_size = 500
        self._min_node_count = 30
        self._max_node_count = 50
        self._min_connect_count = 2
        self._max_connect_count = 5

        # main variable   --------------------------------------------------
        self.nodes_list = []
        self.connect_list = []
        self.system_time = 0
        self.id_counter = 1

        # Control variables  (variables for collecting statistics) --------------------------------------------------
        self.reports_list = []
        self.lost_list = []
        self.best_path = []
        self.gloabal_statistics_list = []
        self.gloabal_history_list = []

        # debug variables   --------------------------------------------------
        self.debug_nodes_current_state_list = []
        self.flag_debug = False
        self.investigation_flag = False
        self.investigation_queue = []

        # jamm parameters   --------------------------------------------------
        self.jamm_x = 700
        self.jamm_y = 800
        self.jamm_power = 30000  # mkwat

    # ---------------------------------------------------------------------------------------------------------------
    def create_net(self, net_type = 'LBZ'):
        self.gloabal_statistics_list = []
        self.lost_list = []
        self.best_path = []
        self.reports_list = []
        self.investigation_queue = []
        self.debug_nodes_current_state_list = []

        if net_type == 'LBZ':
            self.create_net_LBZ_v2()
        else:
            self.create_net_default()

    # ---------------------------------------------------------------------------------------------------------------
    def create_net_LBZ(self):
        nodes_count = random.randint(self._min_node_count, self._max_node_count)
        for xx in range(nodes_count):
            new_node = adhoc_nodes.AdHocNode(
                node_id=xx,
                y=random.randint(0, self._plase_y_size),
                x=random.randint(0, self._plase_x_size),
                model_air=self,
            )
            new_node.__conection_count = random.randint(
                self._min_connect_count, self._max_connect_count
            )
            self.nodes_list.append(new_node)

        neighborhood_threhold = (((self._plase_y_size**2) + (self._plase_x_size**2))**0.5)/10 #Pavlo said divide by 10

    #         check neighborhood
        for nod in self.nodes_list:
            neighborhood_list = self._find_neighborhood( nod.position_x, nod.position_y, neighborhood_threhold)
            while len(neighborhood_list) < 3:
                nod.position_x=random.randint(0, self._plase_x_size)
                nod.position_y=random.randint(0, self._plase_y_size)
                neighborhood_list = self._find_neighborhood(nod.position_x, nod.position_y, neighborhood_threhold)
        #  add connects
        for nod in self.nodes_list:
            neighborhood_list = self._find_neighborhood( nod.position_x, nod.position_y, neighborhood_threhold)
            for neighborhood_node in neighborhood_list:
                if not nod.check_connect_to(neighborhood_node):
                    pneighborhood_node = self.get_node_by_id(neighborhood_node)
                    new_connect = adhoc_nodes.AdHocConnect(nod, pneighborhood_node)
                    nod.connect_list.append(new_connect)
                    pneighborhood_node.connect_list.append(new_connect)
                    self.connect_list.append(new_connect)

        self.set_jamm_to_connects()

    # ---------------------------------------------------------------------------------------------------------------
    def create_net_LBZ_v2(self):
        nodes_count = random.randint(self._min_node_count, self._max_node_count)
        row_count = 3
        column_count = nodes_count/row_count
        step_x = self._plase_x_size / (column_count)
        step_y = self._plase_y_size / row_count
        sigma_x = int(1.1*step_x)
        sigma_y = int(1.1 * step_y)
        for xx in range(nodes_count):
            pos_x = int((xx // row_count) * step_x)
            pos_y =int( (xx % row_count) * step_y)
            pos_x = random.randint(max([0,pos_x-sigma_x]), min([self._plase_x_size-1,pos_x+sigma_x]))
            pos_y = random.randint(max([0, pos_y - sigma_y]), min([self._plase_y_size - 1, pos_y + sigma_y]))

            new_node = adhoc_nodes.AdHocNode(
                node_id=xx,
                y=pos_y,
                x=pos_x,
                model_air=self,
            )
            new_node.__conection_count = random.randint(
                self._min_connect_count, self._max_connect_count
            )
            self.nodes_list.append(new_node)

        neighborhood_threhold = (((self._plase_y_size**2) + (self._plase_x_size**2))**0.5)/7

    #         check neighborhood
        for nod in self.nodes_list:
            neighborhood_list = self._find_neighborhood( nod.position_x, nod.position_y, neighborhood_threhold)
            while len(neighborhood_list) < 3:
                nod.position_x=random.randint(0, self._plase_x_size)
                nod.position_y=random.randint(0, self._plase_y_size)
                neighborhood_list = self._find_neighborhood(nod.position_x, nod.position_y, neighborhood_threhold)
        #  add connects
        for nod in self.nodes_list:
            neighborhood_list = self._find_neighborhood( nod.position_x, nod.position_y, neighborhood_threhold)
            if len(neighborhood_list) > 3:
                neighborhood_list = random.sample(neighborhood_list,3)
            for neighborhood_node in neighborhood_list:
                if not nod.check_connect_to(neighborhood_node):
                    pneighborhood_node = self.get_node_by_id(neighborhood_node)
                    new_connect = adhoc_nodes.AdHocConnect(nod, pneighborhood_node)
                    nod.connect_list.append(new_connect)
                    pneighborhood_node.connect_list.append(new_connect)
                    self.connect_list.append(new_connect)

        self.set_jamm_to_connects()

    # ---------------------------------------------------------------------------------------------------------------
    def _find_neighborhood(self, x, y, max_dist):
        tmp_dist = [[nn.node_id, ( ((nn.position_x - x) ** 2) + ((nn.position_y -y) ** 2))  ** 0.5 ]
            for nn in self.nodes_list ]
        rez_list = [nn[0] for nn in  tmp_dist if ((nn[1] < max_dist) and (nn[1] > 1))]
        return rez_list

    # ---------------------------------------------------------------------------------------------------------------
    def create_net_default(self):
        nodes_count = random.randint(self._min_node_count, self._max_node_count)
        for xx in range(nodes_count):
            new_node = adhoc_nodes.AdHocNode(
                node_id=xx,
                y=random.randint(0, self._plase_y_size),
                x=random.randint(0, self._plase_x_size),
                model_air=self,
            )
            new_node.__conection_count = random.randint(
                self._min_connect_count, self._max_connect_count
            )
            self.nodes_list.append(new_node)

        # -add connect--------------
        for xx in range(nodes_count):
            cur_node = self.nodes_list[xx]
            tmp_dist = [
                [
                    nn.node_id,
                    (
                        ((nn.position_x - cur_node.position_x) ** 2)
                        + ((nn.position_y - cur_node.position_y) ** 2)
                    )
                    ** 0.5,
                ]
                for nn in self.nodes_list
            ]
            tmp_dist.sort(key=lambda x: x[1])

            if len(cur_node.connect_list) < cur_node.__conection_count:
                for ind in range(1, len(tmp_dist)):
                    other_node = self.get_node_by_id(tmp_dist[ind][0])
                    if len(other_node.connect_list) < other_node.__conection_count:
                        if random.randint(0, 100) < 80:
                            new_connect = adhoc_nodes.AdHocConnect(cur_node, other_node)
                            cur_node.connect_list.append(new_connect)
                            other_node.connect_list.append(new_connect)
                            self.connect_list.append(new_connect)
                            if len(cur_node.connect_list) >= cur_node.__conection_count:
                                break
        self.set_jamm_to_connects()

    # ---------------------------------------------------------------------------------------------------------------
    @staticmethod
    def noise_spread(distance, source_noice):
        return source_noice / ((distance / 10) ** 2)

    # ---------------------------------------------------------------------------------------------------------------
    def add_data_pack_for_investigation(self, pack: DataPack):
        if self.investigation_flag:
            return
        self.investigation_queue.append(pack)
        return

    def check_activity(self):
        # for node in self.nodes_list:
        #     for cur_pack in node.queue_to_send:
        #         if isinstance(cur_pack, DataPack):
        #             print(str(cur_pack.current_node_id) + '   ' + str(cur_pack.path))

        for node in self.nodes_list:
            for cur_pack in node.queue_to_send:
                if isinstance(cur_pack, DataPack):
                    return True
        return False

    # ---------------------------------------------------------------------------------------------------------------
    def __copy_nodes_param(self):
        backup_nodes = []
        for node in self.nodes_list:
            backup_nodes.append(
                {
                    "id": node.node_id,
                    "position_x": node.position_x,
                    "position_y": node.position_y,
                }
            )
        return backup_nodes

        # ---------------------------------------------------------------------------------------------------------------

    def __nodes_from_backup(self, backup_nodes):
        self.nodes_list = []
        for node in backup_nodes:
            new_node = adhoc_nodes.AdHocNode(
                node_id=node["id"],
                y=node["position_y"],
                x=node["position_x"],
                model_air=self,
            )
            self.nodes_list.append(new_node)

    # ---------------------------------------------------------------------------------------------------------------
    def __copy_connects_param(self):
        backup_connect = []
        for connect in self.connect_list:
            backup_connect.append(
                {
                    "source": connect.nodes_pointers[0].node_id,
                    "target": connect.nodes_pointers[1].node_id,
                    "default_bit_rate_error": connect.default_bit_rate_error,
                    "default_byte_per_tik": connect.default_byte_per_tik,
                    "signal_value": connect.signal_value,
                }
            )
        return backup_connect

    # ---------------------------------------------------------------------------------------------------------------
    def __connects_from_backup(self, backup_connect):
        self.connect_list = []
        for connect in backup_connect:
            cur_node = self.get_node_by_id(connect["source"])
            other_node = self.get_node_by_id(connect["target"])
            new_connect = adhoc_nodes.AdHocConnect(cur_node, other_node)
            new_connect.default_bit_rate_error = connect["default_bit_rate_error"]
            new_connect.default_byte_per_tik = connect["default_byte_per_tik"]
            new_connect.signal_value = connect["signal_value"]
            self.connect_list.append(new_connect)
        self.set_jamm_to_connects()

    # ---------------------------------------------------------------------------------------------------------------
    def run_investigation(self):
        if self.investigation_flag or len(self.investigation_queue) == 0:
            return
        self.investigation_flag = True
        self.best_path = []
        backup_nodes = self.__copy_nodes_param()
        backup_connect = self.__copy_connects_param()
        investigation_time = self.system_time

        for q in self.investigation_queue:
            self.__nodes_from_backup(backup_nodes)
            self.__connects_from_backup(backup_connect)
            self.system_time = investigation_time
            source_node = self.get_node_by_id(q.path[0])
            source_node._add_to_queue_to_send(q)
            # q.set_current_node(q.path[0])
            # source_node.queue_to_send.append(q)
            while self.check_activity():
                self.turn_one_tik()
            out_line = (
                "\rinvestigation "
                + str(self.investigation_queue.index(q) + 1)
                + " // "
                + str(len(self.investigation_queue))
            )
            print(out_line, end="", flush=True)
        self.investigation_queue = []
        self.__nodes_from_backup(backup_nodes)
        self.__connects_from_backup(backup_connect)
        self.system_time = investigation_time
        self.collect_gloabal_statistics()
        self.investigation_flag = False

    # ---------------------------------------------------------------------------------------------------------------
    def set_jamm_to_connects(self):
        for connect in self.connect_list:
            # for node in connect.nodes_pointers:
            #     distance = ((node.position_y - self.jamm_y) ** 2 +
            #                 (node.position_x - self.jamm_x) ** 2) ** 0.5
            #     print(f"{distance:.2f}")
            noice_value = [
                self.noise_spread(
                    distance=(
                        (node.position_y - self.jamm_y) ** 2
                        + (node.position_x - self.jamm_x) ** 2
                    )
                    ** 0.5,
                    source_noice=self.jamm_power,
                )
                for node in connect.nodes_pointers
            ]
            noice_value.append(2.0)
            connect.noice_value = max(noice_value)
            connect.calculate_error_probability()

    # ---------------------------------------------------------------------------------------------------------------
    def get_current_time(self):
        return self.system_time

    # ---------------------------------------------------------------------------------------------------------------
    def get_new_pack_id(self):
        self.id_counter += 1
        return self.id_counter

    # ---------------------------------------------------------------------------------------------------------------
    def get_node_by_id(self, node_id):
        for node in self.nodes_list:
            if node.node_id == node_id:
                return node
        return None

    # ---------------------------------------------------------------------------------------------------------------
    def send_message(
        self, source_node_id: int, destination_node_id: int, message_length: int
    ):
        node = self.get_node_by_id(source_node_id)
        node.send_message_to(destination_node_id, message_length)

    # ---------------------------------------------------------------------------------------------------------------
    def add_report(self, source_node, target_node, path, full_time, pack_size):
        self.reports_list.append(
            [
                self.get_current_time(),
                source_node,
                target_node,
                pack_size,
                str(path),
                len(path),
                full_time,

            ]
        )
        self.best_path.append([path, full_time , pack_size])
        return

    def add_lost(self, source_node, target_node, path, break_node, pack_type):
        self.lost_list.append(
            [
                self.get_current_time(),
                source_node,
                target_node,
                str(path),
                len(path),
                break_node,
                pack_type,
            ]
        )
        return

    # ---------------------------------------------------------------------------------------------------------------
    def show_net(self):
        output_image = numpy.zeros(
            [self._plase_y_size, self._plase_x_size, 3], dtype=numpy.uint8
        )
        # output jamm
        for x in range(1400, 100, -10):
            col_int = min(int(255 / ((x / 350) ** 2)), 255)
            cv2.circle(
                output_image, (self.jamm_x, self.jamm_y), x, (30, 0, col_int), -1
            )

        # output ad_hoc net
        for connect in self.connect_list:
            n1 = connect.nodes_pointers[0]
            n2 = connect.nodes_pointers[1]
            cv2.line(
                output_image,
                (n1.position_x, n1.position_y),
                (n2.position_x, n2.position_y),
                (120, 0, 0),
                2,
            )

        for node in self.nodes_list:
            cv2.circle(
                output_image, (node.position_x, node.position_y), 5, (0, 120, 0), 2
            )
            cv2.putText(
                output_image,
                str(node.node_id),
                (node.position_x - 2, node.position_y - 2),
                cv2.FONT_HERSHEY_SIMPLEX,
                0.5,
                (200, 200, 200),
                1,
                cv2.LINE_AA,
            )

        if len(self.best_path) > 0:
            self.best_path.sort(key=lambda x: x[1])
            for vv in range(min([len(self.best_path), 5])):
                path = self.best_path[vv][0]
                for nn in range(1, len(path)):
                    n1 = self.get_node_by_id(path[nn - 1])
                    n2 = self.get_node_by_id(path[nn])
                    cv2.line(
                        output_image,
                        (n1.position_x, n1.position_y),
                        (n2.position_x, n2.position_y),
                        (0, 250, 0),
                        4,
                    )
            self.best_path.sort(key=lambda x: len(x[0]))
            for vv in range(min([len(self.best_path), 5])):
                path = self.best_path[vv][0]
                for nn in range(1, len(path)):
                    n1 = self.get_node_by_id(path[nn - 1])
                    n2 = self.get_node_by_id(path[nn])
                    cv2.line(
                        output_image,
                        (n1.position_x, n1.position_y),
                        (n2.position_x, n2.position_y),
                        (200, 0, 200),
                        2,
                    )

        cv2.imshow("image", output_image)
        cv2.waitKey(0)
        cv2.destroyAllWindows()

    # ---------------------------------------------------------------------------------------------------------------
    def collect_debug_information(self):
        self.debug_nodes_current_state_list.append([str(self.get_current_time())])
        rez = ["queue_to_send"]
        for node in self.nodes_list:
            state = ""
            # if node.node_id == 11:
            #     print("Test")
            for q in node.queue_to_send:
                if q is RReqPack or q is DataPack:
                    state = (
                        state
                        + q.type
                        + ": trgt-"
                        + str(q.destination_node_id)
                        + ":"
                        + str(q.path)
                        + ";\n"
                    )
                else:
                    state = state + q.type + ":" + str(q.path) + ";\n"
            rez.append(state)
        self.debug_nodes_current_state_list.append(rez)

        rez = ["queue_wait"]
        for node in self.nodes_list:
            state = ""
            for q in node.queue_wait:
                state = state + str(q.get_next_hop()) + ";"
            rez.append(state)
        self.debug_nodes_current_state_list.append(rez)

    # ---------------------------------------------------------------------------------------------------------------
    def collect_gloabal_statistics(self):
        rez_row_header = ['hop count']
        rez_row_average_time = ['average delay']
        rez_row_best_time = ['min delay']
        rez_row_path_count = ['path count']
        rez_row_path_better_then_DSR = ['number of paths that are better than choosing the standard DSR']

        best_time = 0
        hop_for_bnest_time = 0
        time_for_min_hop = 0
        if len(self.best_path) > 0:
            min_path = min([len(x[0]) for x in self.best_path])
            hop_for_bnest_time = min_path
            select_path_time = [x[1] for x in self.best_path if len(x[0]) == min_path]
            time_for_min_hop = sum(select_path_time) / len(select_path_time)

            best_time = time_for_min_hop
            for hop in range(min_path, min_path + 15):
                select_path_time = [x[1] for x in self.best_path if len(x[0])==hop]
                b_dsr = [x for x in select_path_time if x < time_for_min_hop]
                rez_row_path_better_then_DSR.append(len(b_dsr))
                rez_row_path_count.append(len(select_path_time))
                rez_row_header.append(hop)
                if len(select_path_time) > 0:
                    min_time = min(select_path_time)
                    rez_row_best_time.append(min_time)
                    rez_row_average_time.append(sum(select_path_time) / len(select_path_time))
                    if best_time > min_time:
                        best_time = min_time
                        hop_for_bnest_time = hop
                else:
                    rez_row_average_time.append('')
                    rez_row_best_time.append('')


            self.gloabal_statistics_list.append([len(self.nodes_list),len(self.connect_list), self.jamm_power,
                                                 self.best_path[0][2],
                                                 min_path,time_for_min_hop,
                                                 hop_for_bnest_time,best_time])
            self.gloabal_history_list.append(rez_row_header)
            self.gloabal_history_list.append(rez_row_average_time)
            self.gloabal_history_list.append(rez_row_best_time)
            self.gloabal_history_list.append(rez_row_path_count)
            self.gloabal_history_list.append(rez_row_path_better_then_DSR)

    # ---------------------------------------------------------------------------------------------------------------
    def save_debug_information(self, file_name):
        wb = openpyxl.Workbook()
        curr_sheet = wb.worksheets[0]
        curr_sheet.title = "State"
        rez = ["nodes_id"]
        for node in self.nodes_list:
            rez.append(str(node.node_id))
        curr_sheet.append(rez)
        for dbg in self.debug_nodes_current_state_list:
            curr_sheet.append(dbg)

        try:
            wb.save(file_name)
        except Exception as e:
            print(f"\nFailed to save '{file_name}', cause: {e}")

    # ---------------------------------------------------------------------------------------------------------------
    def save_statistics_information(self, file_name):
        wb = openpyxl.Workbook()
        curr_sheet = wb.worksheets[0]
        curr_sheet.title = "Pack"
        curr_sheet.append(
            [
                "current time",
                "source node",
                "destination node",
                "path",
                "hop",
                "delay time",
            ]
        )
        for row in self.reports_list:
            curr_sheet.append(row)

        curr_sheet = wb.create_sheet("lost")
        curr_sheet.append(
            [
                "current time",
                "source node",
                "destination node",
                "path",
                "hop",
                "break",
                "pack type",
            ]
        )
        for row in self.lost_list:
            curr_sheet.append(row)

        curr_sheet = wb.create_sheet("state")
        curr_sheet.append(
            [
                "first node",
                "second node",
                "signal_value",
                "noice_value",
                "byte_per_tik",
            ]
        )
        for connect in self.connect_list:
            curr_sheet.append(
                [
                    str(connect.nodes_pointers[0].node_id),
                    str(connect.nodes_pointers[1].node_id),
                    str(connect.signal_value),
                    f"{connect.noice_value:.2f}",
                    f"{connect.current_byte_per_tik:.1f}",
                ]
            )

        try:
            wb.save(file_name)
        except Exception as e:
            print(f"\nFailed to save '{file_name}', cause: {e}")

    # ---------------------------------------------------------------------------------------------------------------
    def save_net_to_file(self, file_name):
        wb = openpyxl.Workbook()
        curr_sheet = wb.worksheets[0]
        curr_sheet.title = "Nodes"
        curr_sheet.append(["node_id", "position_y", "position_x"])
        for node in self.nodes_list:
            curr_sheet.append(
                [str(node.node_id), str(node.position_y), str(node.position_x)]
            )

        curr_sheet = wb.create_sheet("Connect")
        curr_sheet.append(
            [
                "first node",
                "second node",
                "default_bit_rate_error",
                "default_byte_per_tik",
            ]
        )
        for connect in self.connect_list:
            curr_sheet.append(
                [
                    str(connect.nodes_pointers[0].node_id),
                    str(connect.nodes_pointers[1].node_id),
                    str(connect.default_bit_rate_error),
                    str(connect.default_byte_per_tik),
                ]
            )

        try:
            wb.save(file_name)
        except Exception as e:
            print(f"\nFailed to save '{file_name}', cause: {e}")

    # ---------------------------------------------------------------------------------------------------------------
    def load_net_from_file(self, file_name):
        self.nodes_list = []
        self.connect_list = []
        self.system_time = 0
        self.id_counter = 1
        self.reports_list = []
        self.debug_nodes_current_state_list = []

        wb = openpyxl.load_workbook(file_name, data_only=True)
        # load nodes
        curr_sheet = wb["Nodes"]
        for row in range(2, curr_sheet.max_row + 1):
            new_node = adhoc_nodes.AdHocNode(
                node_id=int(curr_sheet.cell(row=row, column=1).value),
                y=int(curr_sheet.cell(row=row, column=2).value),
                x=int(curr_sheet.cell(row=row, column=3).value),
                model_air=self,
            )
            self.nodes_list.append(new_node)
        # load Connect
        curr_sheet = wb["Connect"]
        for row in range(2, curr_sheet.max_row + 1):
            n1 = int(curr_sheet.cell(row=row, column=1).value)
            n2 = int(curr_sheet.cell(row=row, column=2).value)
            cur_node = self.get_node_by_id(n1)
            other_node = self.get_node_by_id(n2)
            new_connect = adhoc_nodes.AdHocConnect(cur_node, other_node)
            new_connect.default_bit_rate_error = float(
                curr_sheet.cell(row=row, column=3).value
            )
            new_connect.default_byte_per_tik = float(
                curr_sheet.cell(row=row, column=4).value
            )
            cur_node.connect_list.append(new_connect)
            other_node.connect_list.append(new_connect)
            self.connect_list.append(new_connect)
        self.set_jamm_to_connects()

    # ---------------------------------------------------------------------------------------------------------------
    def turn_one_tik(self):
        self.system_time += 1
        for connect in self.connect_list:
            connect.send_one_tik_part()

        for node in self.nodes_list:
            node.receive_tik()
        # print(self.system_time)

    # ---------------------------------------------------------------------------------------------------------------
    def _get_node_from(self,xmin, xmax, ymin, ymax,):
        while(True):
            select_node = [x for x in self.nodes_list if x.position_x>xmin and x.position_x< xmax]
            if len(select_node) > 0:
                return random.choice(select_node)
            else:
                xmin = max([0,  xmin-30])
                xmax = min([self._plase_x_size-1, xmax+30])


# ---------------------------------------------------------------------------------------------------------------
def save_gloabal_statistics( file_name, gloabal_history_list, gloabal_statistics_list):
    wb = openpyxl.Workbook()
    curr_sheet = wb.worksheets[0]
    curr_sheet.title = "gloabal_history"
    for dbg in gloabal_history_list:
        curr_sheet.append(dbg)
    curr_sheet = wb.create_sheet("gloabal_statistics")
    curr_sheet.append(['nodes count', 'connects count', 'jamm power','size', 'min hop',
                      'DSR time', 'hop for best path', 'best time'])
    for dbg in gloabal_statistics_list:
        curr_sheet.append(dbg)
    try:
        wb.save(file_name)
    except Exception as e:
        print(f"\nFailed to save '{file_name}', cause: {e}")


# ---------------------------------------------------------------------------------------------------------------
#  Run collect statistic for compare DSR and NN
# ----------------------------------------------------------------------------------------------------------------
# if __name__ == "__main__":
#
#     # ad_hoc.run_debugs()
#     run_research_pack_size()


    # ad_hoc.create_net()
    # current_directory = os.getcwd()
    # # ad_hoc.save_net_to_file(current_directory + r'\AdHoc_Net_test.xlsx')
    # ad_hoc.load_net_from_file(file_name=current_directory + r"\AdHoc_Net_test.xlsx")
    # ad_hoc.show_net()
    # ad_hoc.send_message(source_node_id=37, destination_node_id=23, message_length=1000)
    # for i in range(200):
    #     ad_hoc.collect_debug_information()
    #     ad_hoc.turn_one_tik()
    #     ad_hoc.run_investigation()
    # ad_hoc.show_net()
    #
    # ad_hoc.save_debug_information(file_name=current_directory + r"\debug.xlsx")
    # ad_hoc.save_statistics_information(  file_name=current_directory + r"\statistics.xlsx")
    # ad_hoc.save_gloabal_statistics(file_name=current_directory + r"\global.xlsx")
